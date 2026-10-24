"""Adapters around Ding's original processors; no cwd changes or source writes.

The four legacy_processors modules are byte-for-byte copies of the supplied
scripts. Only pure helpers are called here, never their interactive main().
"""
from contextlib import contextmanager
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import re
import sys
import tempfile
import threading
import zipfile
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from ec_export import validate_output_dir, unused_path
from ec_cdi_metadata import suggest_cdi_mass
from legacy_processors import electrochem as ec, cp_cycle as cp, cv_cycle as cv, cdi

BASE = Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).parent
TEMPLATE = BASE / "templates" / "电容贡献率模板.xlsx"
_LOCK = threading.Lock()
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


@contextmanager
def parameters(module, **values):
    old = {k: getattr(module, k) for k in values}
    try:
        for k, v in values.items():
            setattr(module, k, v)
        yield
    finally:
        for k, v in old.items():
            setattr(module, k, v)


def _positive(value, name):
    value = float(value)
    if not math.isfinite(value) or value <= 0:
        raise ValueError(name + "必须是大于 0 的数值。")
    return value


def _label(path):
    # Preserve the old short numeric headings, including decimal conditions.
    match = re.search(r"\d+(?:\.\d+)?", path.stem)
    return format(float(match.group()), "g") if match else path.stem


def _ordered(paths):
    labels = [_label(p) for p in paths]
    if len(set(labels)) != len(labels):
        raise ValueError("本批次有重复的条件名称，旧版宽表会覆盖同名列。请按同一样品、同一质量分批处理。")
    return sorted(paths, key=lambda p: (0, float(_label(p))) if re.fullmatch(r"\d+(?:\.\d+)?", _label(p)) else (1, _label(p)))


def electrochemistry(paths, folder, mass_mg=24, time_step=0.1, progress=print):
    mass = _positive(mass_mg, "质量") / 1000
    time_step = _positive(time_step, "GCD 采样间隔")
    groups = {kind: [] for kind in ("CV", "GCD")}
    for path in paths:
        kind = ec.detect_type(ec.read_text(path))
        if kind not in groups:
            raise ValueError(f"{path.name} 不是旧脚本支持的 CV/GCD TXT。")
        groups[kind].append(path)
    made = []
    if groups["CV"]:
        info = []
        for path in _ordered(groups["CV"]):
            df = ec.extract_cv(path)
            if df.empty:
                raise ValueError(path.name + " 没有有效 CV 数据。")
            info.append((path, df, ec.cv_turn_points(df.Potential.to_numpy())))
        start, end = (2, 3) if min(len(tp)-1 for _, _, tp in info) < 4 else (3, 4)
        out = None
        for path, df, tp in info:
            if len(tp) < end + 1:
                raise ValueError(path.name + " 圈数不足，无法按原脚本选段。")
            seg = df.iloc[tp[start-1][0]:tp[end][0]+1].reset_index(drop=True)
            potential = seg.Potential.to_numpy()
            if out is None:
                out = pd.DataFrame({"Potential": potential})
            elif len(out) != len(potential) or not np.allclose(out.Potential, potential, atol=ec.CV_TOLERANCE):
                raise ValueError(path.name + " 电压网格与其他 CV 文件不一致，不能共用一列电压。请分批处理。")
            out[_label(path)] = seg.Current / mass
        destination = folder / "CV结果.xlsx"
        out.to_excel(destination, index=False)
        made.append(destination)
        progress(f"CV：沿用原脚本第 {start}+{end} 段、电流除以 {mass_mg:g} mg；一列电压、多列电流密度。")
    if groups["GCD"]:
        values, durations, marks = {}, [], {}
        for path in _ordered(groups["GCD"]):
            df = ec.extract_gcd(path)
            y, t = df.Voltage.to_numpy(), df.Time.to_numpy()
            peaks, valleys, events = ec.get_events(y)
            if len(events) < ec.GCD_END_CYCLE:
                raise ValueError(path.name + " 峰谷数量不足，无法按原脚本截取第 2+3 段。")
            segment = y[events[ec.GCD_START_CYCLE-2][0]:events[ec.GCD_END_CYCLE-1][0]+1]
            label = _label(path)
            values[label] = segment
            marks[label] = int(np.argmax(segment))+2
            if len(peaks) and len(valleys) > 1:
                durations.append([int(label) if label.isdigit() else label, t[valleys[1]]-t[peaks[0]]])
            if len(t) > 1 and not np.allclose(np.diff(t), time_step, atol=0.0001):
                progress(path.name + "：仪器时间间隔与设定不完全一致，Time 列仍按原脚本的固定采样间隔生成。")
        out = pd.DataFrame({"Time": np.arange(max(map(len, values.values()))) * time_step})
        for label, value in values.items():
            out[label] = pd.Series(value)
        destination = folder / "GCD结果.xlsx"
        with pd.ExcelWriter(destination, engine="openpyxl") as writer:
            out.to_excel(writer, sheet_name="GCD_Data", index=False)
            pd.DataFrame(durations, columns=["Sample", "Peak_to_Valley_Time"]).to_excel(writer, sheet_name="Duration", index=False)
            sheet = writer.sheets["GCD_Data"]
            for col, label in enumerate(values, 2):
                sheet.cell(marks[label], col).fill = PatternFill("solid", fgColor="FFFF00")
        made.append(destination)
        progress("GCD：保留 GCD_Data、Duration 两张表与黄色峰值标记。")
    return made


def cp_cycles(path, folder, time_step=0.1):
    with parameters(cp, TIME_STEP=_positive(time_step, "采样间隔")):
        voltage = cp.process_raw_txt(str(path))
        if len(voltage) < 5:
            raise ValueError("CP 数据不足。")
        troughs = cp.find_all_troughs(voltage)
        cycles = cp.split_cycles(voltage, troughs)
        if max(map(len, cycles.values())) * len(cycles) > 500_000:
            raise ValueError("CP 宽表超过 50 万个单元格，请分批处理。")
        durations = cp.calc_discharge_time(cycles, voltage, troughs)
        destination = Path(unused_path(folder / f"{path.stem}_处理结果.xlsx"))
        cp.export_xlsx(cp.build_export_df(cycles), durations, str(destination))
        return destination


def cv_cycles(path, folder, mode="peak", progress=print):
    voltage, current = cv.read_and_clean_txt(str(path))
    cycles = cv.split_cycles(voltage, current, cv.detect_points(voltage, mode))
    if not cycles:
        raise ValueError("没有识别出完整 CV 循环，请检查峰/谷分圈选项。")
    if max(len(v) for v, _ in cycles) * len(cycles) > 500_000:
        raise ValueError("CV 循环宽表超过 50 万个单元格，请分批处理。")
    base = np.asarray(cycles[0][0])
    if any(len(v) != len(base) or not np.allclose(v, base, atol=1e-6) for v, _ in cycles[1:]):
        progress("注意：各圈电压网格不完全一致；沿用旧脚本的首圈电压列，作图前请核对。")
    tables = cv.filter_by_last_digit(cv.build_processed_table(cycles))
    destination = Path(unused_path(folder / f"{path.stem}.xlsx"))
    cv.save_xlsx_all(tables, str(destination.with_suffix("")))
    return destination


def _cache_formulas(path, caches):
    """Store known numeric results alongside formulas; Excel can still recalc.

    All other zip entries (styles, merges, chart objects) are kept unchanged.
    This is a cache writer, not an arbitrary formula interpreter.
    """
    temp = path.with_suffix(".cache.tmp")
    try:
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
            for entry in source.infolist():
                raw = source.read(entry.filename)
                if entry.filename in caches:
                    cache = caches[entry.filename]
                    def update_cell(match):
                        opening, body = match.group(1), match.group(2)
                        if "<f" not in body:
                            return match.group(0)
                        coordinate = re.search(r'\br="([^"]+)"', opening).group(1)
                        if coordinate not in cache:
                            raise ValueError("模板含未识别的公式位置：" + entry.filename + "!" + coordinate)
                        value = cache[coordinate]
                        opening = re.sub(r'\s+t="[^"]*"', "", opening)
                        body = re.sub(r"<v(?:\s[^>]*)?>.*?</v>|<v\s*/>", "", body, flags=re.S)
                        if value is None or isinstance(value, str):
                            opening = opening[:-1] + ' t="str">'
                            content = escape("" if value is None else value)
                        else:
                            if not math.isfinite(float(value)):
                                raise ValueError("公式计算产生无效数值：" + coordinate)
                            content = repr(float(value))
                        return opening + body + "<v>" + content + "</v></c>"
                    raw = re.sub(r"(<c\b[^>]*>)(.*?)</c>", update_cell, raw.decode("utf-8"), flags=re.S).encode("utf-8")
                target.writestr(entry, raw)
        temp.replace(path)
    finally:
        if temp.exists():
            temp.unlink()


def _cdi_time_step(times, requested):
    automatic = requested is None or str(requested).strip().lower() == "auto"
    if automatic:
        if len(times) < 2 or not np.isfinite(times).all():
            raise ValueError("无法从完整时间戳识别 CDI 采样间隔，请手动填写秒/点。")
        intervals = np.diff(times)
        if np.any(intervals <= 0):
            raise ValueError("CDI 时间重复或倒退，无法自动识别采样间隔，请核对原始时间。")
        step = float(np.median(intervals))
    else:
        step = _positive(requested, "采样间隔")
    return step, automatic


def cdi_workbook(path, folder, options, progress=print, metadata=None):
    from ec_conductivity import process_files
    mass = _positive(options.get("mass_mg", 24), "质量")
    minutes = _positive(options.get("cycle_minutes", 60), "每圈时长")
    volume_l = _positive(options.get("volume_ml", 50), "溶液体积") / 1000
    base = float(options.get("calibration_base", 8.2236))
    scale = _positive(options.get("calibration_scale", 1.97252), "标定斜率")
    result = process_files([str(path)], bare_unit=options.get("conductivity_unit", "uS/cm"))
    frame = result.data
    if frame.empty or frame.conductivity_uS_cm.isna().any() or frame.source_sheet.nunique() != 1:
        raise ValueError("电导率文件含缺失或无效测量，不能直接按固定行数分圈。请先用单位检查工具核对。")
    values = frame.conductivity_uS_cm.to_numpy(dtype=float)
    times = frame.elapsed_s.to_numpy(dtype=float)
    step, automatic = _cdi_time_step(times, options.get("cdi_step", "auto"))
    count = minutes * 60 / step
    if not math.isfinite(base) or not math.isclose(count, round(count)) or count < 2 or count > 100_000:
        raise ValueError("每圈时长必须能被采样间隔整除，且每圈为 2–10 万个点。")
    count = int(round(count))
    if automatic:
        progress(f"从实测时间自动识别：{step:g} 秒/点；每圈 {minutes:g} 分钟、{count} 点。")
    if np.isnan(times).all():
        times = np.arange(len(values))*step
        progress(f"源文件无有效时间，按设定的 {step:g} 秒/点生成 Time。")
    elif (not np.isfinite(times).all() or np.any(np.diff(times) <= 0)
          or not np.allclose(times, np.arange(len(times))*step, atol=max(1., step*.02), rtol=0)):
        raise ValueError("实测时间间隔与设定不一致，不能按固定行数分圈；请核对采样间隔或使用标准电导率处理。")
    elif not np.allclose(np.diff(times), step, atol=.01):
        progress("时间戳有轻微抖动（累计偏差不超过 1 秒或采样间隔的 2%）；保留实测时间，按原脚本的每圈行数分圈。")
    tail = options.get("cdi_tail", "reference")
    if tail not in {"reference", "complete", "include"}:
        raise ValueError("未知的 CDI 末圈策略。")
    cycle_values = values
    remainder = len(values) % count
    # The supplied 7201-point reference keeps the closing t=36000 sample in
    # raw data while exporting ten 720-point cycles, not an eleventh cycle.
    closing_endpoint = (tail == "reference" and len(values) > count and remainder == 1
                        and math.isclose(times[-1], (len(values)-1)*step,
                                         abs_tol=max(1., step*.02)))
    if tail == "complete" and remainder:
        cycle_values = values[:len(values)-remainder]
        progress(f"末尾 {remainder} 点不足一圈，仅保留在原始数据表。")
    elif closing_endpoint:
        cycle_values = values[:-1]
        progress("整圈结束的最后 1 个时间点保留在原始数据表，不另生成只有 1 点的新圈。")
    elif remainder:
        progress(f"末圈不足 {count} 点：保留已有 {remainder} 点，其余留空。")
    if not len(cycle_values):
        raise ValueError("数据不足一个完整圈，请选择保留不完整末圈，或核对每圈时长。")
    circles = math.ceil(len(cycle_values)/count)
    if count*circles*6 > 1_000_000:
        raise ValueError("CDI 七表超过 100 万单元格，请分批处理。")
    with parameters(cdi, TIME_STEP_SECONDS=step, CYCLE_TOTAL_MINUTES=minutes,
                    ROWS_PER_CYCLE=count, CONDUCTIVITY_TO_CONCENTRATION_BASE=base,
                    CONDUCTIVITY_TO_CONCENTRATION_SCALE=scale, SOLUTION_VOLUME_ML=volume_l):
        template = options.get("cdi_template")
        wb = load_workbook(template) if template else Workbook()
        if not template:
            wb.remove(wb.active)
            # Both reference workbooks use the plain Songti 11-point layout.
            wb._fonts[0] = Font(name="宋体", size=11, color="000000")
        cdi.ensure_sheet_order(wb)
        for sheet in wb:
            cdi.clear_sheet_values(sheet)
        cdi.write_raw_sheet(wb["原始数据"], times, values, mass)
        cdi.write_conductivity_sheet(wb["电导率"], len(cycle_values))
        cdi.write_concentration_sheet(wb["浓度"], circles)
        cdi.write_difference_sheet(wb["浓度差"], circles)
        cdi.write_capacity_sheet(wb["吸附容量"], circles)
        cdi.write_rate_sheet(wb["吸附速率"], circles)
        cdi.write_max_sheet(wb["吸附最大值"], circles)
    # Retain original formulas for valid measurements. Guard only absent tail
    # rows and t=0; these used to create artificial concentrations / #DIV/0!.
    caches = {f"xl/worksheets/sheet{i}.xml": {} for i in range(2, 8)}
    for row in range(2, count+2):
        time_value = float(times[row-2]) if row-2 < len(times) else None
        for sheet_index in range(2, 7):
            caches[f"xl/worksheets/sheet{sheet_index}.xml"][f"A{row}"] = time_value
            if time_value is None:
                wb.worksheets[sheet_index-1].cell(row, 1, f'=IF(原始数据!A{row}="","",原始数据!A{row})')
    for circle in range(circles):
        col = get_column_letter(circle+2)
        values_in_cycle = cycle_values[circle*count:(circle+1)*count]
        concentrations = (values_in_cycle-base)/scale
        differences = concentrations[0]-concentrations
        capacity = differences*volume_l/mass*1000
        for offset in range(count):
            row = offset+2
            valid = offset < len(values_in_cycle)
            t = times[offset] if offset < len(times) else None
            computed = [values_in_cycle[offset], concentrations[offset], differences[offset], capacity[offset],
                        capacity[offset]/t if t else None] if valid else [None]*5
            for index, value in enumerate(computed, 2):
                cell = wb.worksheets[index-1].cell(row, circle+2)
                if index > 2 and not valid:
                    cell.value = f'=IF(电导率!{col}{row}="","",{cell.value[1:]})'
                if index == 6 and valid:
                    cell.value = f'=IF(A{row}=0,"",{cell.value[1:]})'
                caches[f"xl/worksheets/sheet{index}.xml"][f"{col}{row}"] = value
        caches["xl/worksheets/sheet7.xml"][f"{get_column_letter(circle+1)}2"] = float(np.max(capacity))
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True)
    destination = Path(unused_path(folder / f"{path.stem}-处理.xlsx"))
    wb.save(destination)
    wb.close()
    _cache_formulas(destination, caches)
    if metadata is not None:
        metadata.update(mass_mg=mass, sampling_seconds=step, automatic_sampling=automatic,
                        cycle_minutes=minutes, points_per_cycle=count, cycles=circles,
                        raw_points=len(values), cycle_points=len(cycle_values),
                        trailing_points_only_in_raw=len(values)-len(cycle_values),
                        tail_policy=tail, volume_ml=volume_l*1000,
                        calibration_base=base, calibration_scale=scale)
    progress("CDI：保留七表与原公式；仅为零秒速率和末圈缺测补空白，避免除零和虚假吸附值。")
    return destination


def run_legacy(mode, paths, output_dir, options=None, progress=print):
    options = dict(options or {})
    paths = [Path(p).resolve() for p in paths]
    if not paths or len(paths) > 100:
        raise ValueError("请选择 1–100 个同批次文件。")
    for path in paths:
        if not path.is_file() or path.stat().st_size > 50*1024*1024:
            raise ValueError("源文件不存在或超过 50 MB：" + str(path))
    protected = [p.parent for p in paths]
    if options.get("cdi_template"):
        protected.append(Path(options["cdi_template"]).resolve().parent)
    output = Path(validate_output_dir(output_dir, protected))
    output.mkdir(parents=True, exist_ok=True)
    folder = Path(tempfile.mkdtemp(prefix="原样Excel_" + datetime.now().strftime("%Y%m%d_%H%M%S") + "_", dir=output))
    made = []
    resolved = []
    with _LOCK:
        try:
            if mode == "electrochem":
                made = electrochemistry(paths, folder, options.get("mass_mg", 24), options.get("time_step", 0.1), progress)
            else:
                for path in paths:
                    progress("正在处理：" + path.name)
                    if mode == "cp_cycle":
                        made.append(cp_cycles(path, folder, options.get("time_step", 0.1)))
                        progress("注意：CP 放电时长沿用原算法，其候选峰实际上按谷值检测，时长需复核；曲线和尾数分表保持原样。")
                    elif mode == "cv_cycle":
                        made.append(cv_cycles(path, folder, options.get("split_mode", "peak"), progress))
                    elif mode == "cdi":
                        details = {"source": str(path)}
                        made.append(cdi_workbook(path, folder, options, progress, details))
                        resolved.append(details)
                    elif mode == "contribution":
                        from ec_contribution import fill_template
                        made.append(fill_template(path, folder))
                    else:
                        raise ValueError("未知处理类型。")
        except Exception as exc:
            (folder / "未完成.txt").write_text(str(exc) + "\n已产生的文件仅供检查，请勿当作完整结果。", encoding="utf-8")
            raise ValueError(f"{exc}\n本次目录：{folder}") from exc
    manifest = {"mode": mode, "parameters": options, "files": [p.name for p in made],
                "sources": [{"path": str(p), "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in paths]}
    if resolved:
        manifest["resolved_parameters"] = resolved
    (folder / "处理说明.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"folder": str(folder), "files": list(map(str, made))}
