# -*- coding: utf-8 -*-
"""Excel/CSV scientific plotting, independent of calculation and Qt.

Presets are ec-lab implementations using Matplotlib's public Figure API. Layout
ideas (inward ticks, small journal figures, monochrome lines) refer to
SciencePlots, https://github.com/garrettj403/SciencePlots (MIT). No SciencePlots
source, dependency, LaTeX installation, or journal certification is included.
Matplotlib style reference: https://matplotlib.org/stable/users/explain/customizing.html
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import sys
import tempfile
import uuid
import zipfile
from dataclasses import asdict, dataclass, field, replace
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.colors import to_hex, to_rgba
from matplotlib.figure import Figure
from openpyxl import load_workbook

BASE_DIR = Path(sys.executable if getattr(sys, "frozen", False) else __file__).resolve().parent
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ROWS = 100_000
MAX_CELLS = 1_000_000
MAX_SERIES = 200
LIMIT_NOTE = "每个文件最多 50 MB、每表 10 万行、整本 100 万单元格、每图 200 条曲线；超限会停止，不会截断数据。"
STYLE_SOURCE = {
    "implementation": "ec-lab native Matplotlib presets; no external style code copied",
    "reference": "https://github.com/garrettj403/SciencePlots",
    "reference_license": "MIT",
    "matplotlib": "https://matplotlib.org/stable/users/explain/customizing.html",
}
PRESETS = {
    "paper": {"name": "论文简洁", "size": (6.4, 4.5), "font": 11,
              "width": 1.5, "grid": False, "mono": False},
    "mono": {"name": "黑白印刷", "size": (6.4, 4.5), "font": 11,
             "width": 1.5, "grid": False, "mono": True},
    "talk": {"name": "报告展示", "size": (7.2, 4.8), "font": 14,
             "width": 2.1, "grid": True, "mono": False},
}
COLORS = ("#0072B2", "#D55E00", "#009E73", "#CC79A7", "#E69F00",
          "#56B4E9", "#333333", "#999999")
FONT = ["Times New Roman", "Microsoft YaHei", "DejaVu Sans"]
LEGEND_LOCATIONS = ("best", "upper right", "upper left", "lower left", "lower right",
                    "right", "center left", "center right", "lower center", "upper center", "center")
MAX_RASTER_PIXELS = 40_000_000
STYLE_TEMPLATE_SCHEMA = "ec-lab-plot-style"
STYLE_TEMPLATE_VERSION = 1
_STYLE_NUMBERS = {
    "axis_font_size": (1, 96), "tick_font_size": (1, 96), "legend_font_size": (1, 96),
    "width_cm": (2, 60), "height_cm": (2, 60), "linewidth": (0.1, 20), "dpi": (72, 1200),
}
_STYLE_KEYS = set(_STYLE_NUMBERS) | {
    "xlim", "ylim", "legend_visible", "legend_location", "series_labels", "series_colors",
    "palette", "font_family",
}


def _style_number(value, name):
    if isinstance(value, (bool, str)) or not isinstance(value, (int, float, np.integer, np.floating)):
        raise ValueError(f"绘图设置 {name} 必须是有限数值。")
    try:
        value = float(value)
    except (ValueError, OverflowError) as exc:
        raise ValueError(f"绘图设置 {name} 必须是有限数值。") from exc
    if not math.isfinite(value):
        raise ValueError(f"绘图设置 {name} 必须是有限数值。")
    return value


def _style_color(value, name, allow_empty=False):
    if not isinstance(value, str):
        raise ValueError(f"绘图设置 {name} 必须是颜色名称或十六进制颜色。")
    value = value.strip()
    if not value and allow_empty:
        return ""
    # C0-style references depend on global rcParams and are not reproducible.
    if not value or re.fullmatch(r"C\d+", value):
        raise ValueError(f"绘图设置 {name} 的颜色无效：{value!r}。")
    try:
        rgba = to_rgba(value)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"绘图设置 {name} 的颜色无效：{value!r}。") from exc
    return to_hex(rgba, keep_alpha=rgba[3] != 1)


def validate_plot_style(style):
    """Validate declarative appearance only and return an independent sparse dict.

    Missing keys inherit the selected preset. Empty per-curve entries inherit
    the original label or palette. None for legend visibility means <=20 curves.
    No files, data columns, expressions or executable code belong in a style.
    """
    if not isinstance(style, dict):
        raise ValueError("绘图样式必须是设置对象。")
    unknown = set(style) - _STYLE_KEYS
    if unknown:
        raise ValueError("绘图样式含不支持的字段（不能保存数据路径或数据列）：" +
                         ", ".join(sorted(map(str, unknown))))
    clean = {}
    for key, value in style.items():
        if key in _STYLE_NUMBERS:
            number = _style_number(value, key)
            low, high = _STYLE_NUMBERS[key]
            if not low <= number <= high:
                raise ValueError(f"绘图设置 {key} 必须在 {low}～{high} 之间。")
            if key == "dpi" and not number.is_integer():
                raise ValueError("绘图 DPI 必须是整数。")
            clean[key] = int(number) if key == "dpi" else number
        elif key in ("xlim", "ylim"):
            if value is None:
                clean[key] = None
                continue
            if not isinstance(value, (list, tuple)) or len(value) != 2:
                raise ValueError(f"{key} 应为自动范围 null，或 [最小值, 最大值]。")
            limits = [_style_number(item, key) for item in value]
            if limits[0] >= limits[1] or not math.isfinite(limits[1] - limits[0]):
                raise ValueError(f"{key} 的最小值必须小于最大值，且范围必须有限。")
            clean[key] = limits
        elif key == "legend_visible":
            if value is not None and not isinstance(value, bool):
                raise ValueError("图例显示应为 true、false 或 null（自动）。")
            clean[key] = value
        elif key == "legend_location":
            if not isinstance(value, str) or value not in LEGEND_LOCATIONS:
                raise ValueError("不支持的图例位置。")
            clean[key] = value
        elif key == "font_family":
            if not isinstance(value, str) or not value.strip() or len(value) > 100 or any(ord(c) < 32 for c in value):
                raise ValueError("字体名称不能为空，且不能超过 100 个字符或含控制字符。")
            clean[key] = value.strip()
        elif key in ("series_labels", "series_colors", "palette"):
            if not isinstance(value, (list, tuple)) or len(value) > MAX_SERIES:
                raise ValueError(f"{key} 必须是最多 {MAX_SERIES} 项的列表。")
            if key == "palette" and not value:
                raise ValueError("调色盘至少需要一种颜色。")
            if key == "series_labels":
                if any(not isinstance(label, str) or len(label) > 300 or
                       any(ord(c) < 32 for c in label) for label in value):
                    raise ValueError("曲线标签必须是最多 300 个字符且不含控制字符的文字。")
                clean[key] = [label.strip() for label in value]
            else:
                clean[key] = [_style_color(color, key, allow_empty=key == "series_colors") for color in value]
    if {"width_cm", "height_cm", "dpi"}.issubset(clean):
        pixels = math.ceil(clean["width_cm"] / 2.54 * clean["dpi"]) * math.ceil(clean["height_cm"] / 2.54 * clean["dpi"])
        if pixels > MAX_RASTER_PIXELS:
            raise ValueError("图幅尺寸与 DPI 组合超过 4000 万像素，请降低尺寸或 DPI，避免占用过多内存。")
    return clean


def resolve_plot_style(style=None, preset="paper"):
    """Return a complete validated appearance without modifying the input."""
    base = PRESETS.get(preset, PRESETS["paper"])
    effective = {"xlim": None, "ylim": None, "legend_visible": None, "legend_location": "best",
                 "series_labels": [], "series_colors": [],
                 "palette": ["#222222"] if base["mono"] else list(COLORS),
                 "font_family": FONT[0], "axis_font_size": base["font"],
                 "tick_font_size": base["font"] - 1, "legend_font_size": max(7, base["font"] - 2),
                 "width_cm": base["size"][0] * 2.54, "height_cm": base["size"][1] * 2.54,
                 "linewidth": base["width"], "dpi": 600}
    effective.update(validate_plot_style({} if style is None else style))
    return validate_plot_style(effective)


def save_style_template(path, style):
    """Atomically save appearance only; source archives remain read-only."""
    clean = validate_plot_style(style)
    destination = Path(path).expanduser().resolve()
    if destination.suffix.lower() != ".json":
        raise ValueError("绘图样式模板请保存为 .json 文件。")
    from ec_export import validate_output_dir
    validate_output_dir(destination.parent)
    payload = json.dumps({"schema": STYLE_TEMPLATE_SCHEMA, "version": STYLE_TEMPLATE_VERSION,
                          "style": clean}, ensure_ascii=False, indent=2, allow_nan=False)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", suffix=".tmp",
                                         prefix=".plot-style-", dir=destination.parent, delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(payload)
        os.replace(temporary, destination)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()
    return str(destination)


def load_style_template(path):
    """Load a small versioned JSON style without opening its source data."""
    source = Path(path)
    if source.stat().st_size > 256 * 1024:
        raise ValueError("绘图样式模板超过 256 KB，无法导入。")
    try:
        payload = json.loads(source.read_text(encoding="utf-8-sig"))
    except (ValueError, UnicodeError) as exc:
        raise ValueError("绘图样式模板不是有效的 UTF-8 JSON。") from exc
    if not isinstance(payload, dict) or payload.get("schema") != STYLE_TEMPLATE_SCHEMA:
        raise ValueError("这不是 ec-lab 绘图样式模板。")
    if type(payload.get("version")) is not int or payload["version"] != STYLE_TEMPLATE_VERSION:
        raise ValueError("不支持此绘图样式模板版本，请使用版本 1。")
    if set(payload) - {"schema", "version", "style"}:
        raise ValueError("绘图样式模板含不支持的字段；样式模板不能保存数据路径或数据列。")
    return validate_plot_style(payload.get("style"))


@dataclass
class TableBook:
    path: str
    tables: dict[str, pd.DataFrame]
    sha256: str
    notes: list[str] = field(default_factory=list)


@dataclass
class ColumnSeries:
    sheet: str
    x: str
    y: str
    group: str = ""
    filter_column: str = ""
    filter_value: str = ""
    negate_y: bool = False
    mode: str = "line"
    label: str = ""
    labels: dict[str, str] = field(default_factory=dict)


@dataclass
class PlotSpec:
    name: str
    series: list[ColumnSeries]
    xlabel: str = ""
    ylabel: str = ""
    title: str = ""
    equal: bool = False
    style: dict = field(default_factory=dict)


def _hash(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _columns(values):
    result, used = [], set()
    for index, value in enumerate(values):
        name = str(value).strip() if value is not None else ""
        name = name or f"未命名列{index + 1}"
        original, number = name, 2
        while name in used:
            name = f"{original}_{number}"
            number += 1
        used.add(name)
        result.append(name)
    return result


def load_table_book(path: str) -> TableBook:
    """Read source only; enforce bounds before materializing large workbooks."""
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise ValueError("找不到数据文件，请重新选择。")
    if source.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("文件超过 50 MB。请按样品拆成较小的表格再导入。")
    suffix = source.suffix.lower()
    tables, notes = {}, []
    if suffix in (".xlsx", ".xlsm"):
        with zipfile.ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError("Excel 解压后超过 200 MB，请拆分文件。")
        workbook = load_workbook(source, read_only=True, data_only=True)
        cell_count = 0
        try:
            for sheet in workbook:
                if sheet.title.startswith("_ec_lab"):
                    continue
                if (sheet.max_column or 0) > 500 and (sheet.title in ("全量CP数据", "原始处理数据") or re.fullmatch(r"尾[号数]\d圈?", sheet.title)):
                    notes.append(f"「{sheet.title}」超过 500 列，已跳过；请使用较小的尾号分表。")
                    continue
                if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 500:
                    raise ValueError(f"工作表「{sheet.title}」超过 10 万行或 500 列，请拆分。")
                iterator = sheet.iter_rows(values_only=True)
                header = next(iterator, ())
                if not header or not any(value is not None for value in header):
                    continue
                header = list(header)
                while header and header[-1] is None:
                    header.pop()
                rows = []
                for row in iterator:
                    values = list(row[:len(header)])
                    rows.append(values)
                    cell_count += len(header)
                    if len(rows) > MAX_ROWS or cell_count > MAX_CELLS:
                        raise ValueError("表格超过 10 万行或整本超过 100 万单元格，请拆分后导入。")
                if rows:
                    tables[sheet.title] = pd.DataFrame(rows, columns=_columns(header))
        finally:
            workbook.close()
        notes.append("Excel 公式只读取已保存的计算结果；未保存的公式缓存会成为空值，请先在 Excel 中重算并保存。")
    elif suffix in (".csv", ".tsv", ".txt"):
        raw = source.read_bytes()
        decoded = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                decoded = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise ValueError("文本编码无法识别，请另存为 UTF-8 CSV。")
        try:
            dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t")
            separator = dialect.delimiter
        except csv.Error:
            separator = "\t" if suffix == ".tsv" else ","
        frame = pd.read_csv(io.StringIO(decoded), sep=separator, nrows=MAX_ROWS + 1, skip_blank_lines=False)
        if len(frame) > MAX_ROWS or frame.size > MAX_CELLS or len(frame.columns) > 500:
            raise ValueError("文本超过 10 万行、500 列或 100 万单元格，请拆分后导入。")
        frame.columns = _columns(frame.columns)
        tables[source.stem] = frame
    else:
        raise ValueError("请使用 .xlsx、.xlsm、.csv 或 .tsv；旧版 .xls 请先另存为 .xlsx。")
    if not tables:
        raise ValueError("没有读到数据。请把列标题放在第一行，数据从第二行开始。")
    if "吸附最大值" in tables:
        frame = tables["吸附最大值"]
        if len(frame) and all(re.fullmatch(r"\d+(?:\.0)?", c) for c in frame.columns):
            tables["吸附最大值"] = pd.DataFrame({"圈号": [float(c) for c in frame.columns],
                "最大吸附容量": list(frame.iloc[0])})
    return TableBook(str(source), tables, _hash(source), notes)


def numeric_columns(frame):
    result = []
    for name in frame.columns:
        values = frame[name].dropna()
        if len(values) and pd.to_numeric(values, errors="coerce").notna().all():
            result.append(str(name))
    return result


def _has(book, sheet, *columns):
    return sheet in book.tables and set(columns).issubset(book.tables[sheet].columns)


def _normalize_legacy_summary(book):
    """Recreate declared legacy condition columns, including recipe replay."""
    if not _has(book, "summary", "condition", "type"):
        return
    summary = book.tables["summary"]
    for kind, column in (("CV", "scan_rate_mVs"), ("GCD", "current_density_A_g")):
        if column not in summary:
            matches = summary["type"].astype(str).str.upper().eq(kind)
            summary[column] = pd.to_numeric(summary["condition"].astype(str)
                .str.extract(r"^\s*([-+]?\d*\.?\d+(?:[eE][-+]?\d+)?)", expand=False), errors="coerce").where(matches)


def _summary_mode(book, kind, column):
    """Repeated conditions cannot establish a material's connected rate curve."""
    summary = book.tables["summary"]
    matches = summary["type"].astype(str).str.upper().eq(kind)
    conditions = pd.to_numeric(summary.loc[matches, column], errors="coerce").dropna()
    if conditions.duplicated().any():
        note = f"{kind} 存在重复倍率条件，无法可靠区分材料；比电容图使用散点，避免把不同样品连成曲线。"
        if note not in book.notes:
            book.notes.append(note)
        return "scatter"
    return "line_points"


def automatic_plots(book: TableBook) -> list[PlotSpec]:
    """Recognize ec-lab v1/v2 sheets; never rerun scientific calculations."""
    _normalize_legacy_summary(book)
    plots = []
    labels = {}
    if _has(book, "summary", "file", "condition"):
        summary = book.tables["summary"]
        counts = summary["condition"].astype(str).value_counts()
        for _, row in summary.iterrows():
            condition, name = str(row["condition"]), str(row["file"])
            labels[name] = condition if counts.get(condition, 0) == 1 else f"{condition} · {name}"
    if _has(book, "CV", "file", "Potential_V", "CurrentDensity_A_g"):
        plots.append(PlotSpec("CV 叠加图", [ColumnSeries("CV", "Potential_V",
                      "CurrentDensity_A_g", "file", labels=labels)], "Potential (V)", "Current density (A/g)"))
    if _has(book, "GCD", "file", "time_s", "Potential_V"):
        plots.append(PlotSpec("GCD 充放电曲线", [ColumnSeries("GCD", "time_s",
                      "Potential_V", "file", labels=labels)], "Time (s)", "Potential (V)"))
    if _has(book, "summary", "type", "capacitance_F_g"):
        summary = book.tables["summary"]
        for kind, column, xlabel in (("CV", "scan_rate_mVs", "Scan rate (mV/s)"),
                                     ("GCD", "current_density_A_g", "Current density (A/g)")):
            matches = summary["type"].astype(str).str.upper().eq(kind)
            if not matches.any():
                continue
            if column in summary:
                mode = _summary_mode(book, kind, column)
                plots.append(PlotSpec(f"{kind} 比电容图", [ColumnSeries("summary", column,
                    "capacitance_F_g", filter_column="type", filter_value=kind,
                    mode=mode, label=kind)], xlabel, "Specific capacitance (F/g)"))
    if _has(book, "EIS", "file", "Z_re_ohm", "Z_im_ohm"):
        series = [ColumnSeries("EIS", "Z_re_ohm", "Z_im_ohm", "file", negate_y=True, mode="scatter")]
        if _has(book, "EIS_fit_curve", "file", "Z_re_fit_ohm", "Z_im_fit_ohm"):
            series.append(ColumnSeries("EIS_fit_curve", "Z_re_fit_ohm", "Z_im_fit_ohm", "file",
                                       negate_y=True, label="拟合"))
        plots.append(PlotSpec("EIS Nyquist 图", series, "Z′ (Ω)", "−Z″ (Ω)", equal=True))
    for sheet, ylabel, name in (("cycle_capacitance", "Specific capacitance (F/g)", "循环比电容"),
                                 ("cycle_retention", "Retention (%)", "循环保持率")):
        if sheet in book.tables:
            frame = book.tables[sheet]
            x = "cycle" if "cycle" in frame else str(frame.columns[0])
            ys = [column for column in numeric_columns(frame) if column != x]
            if ys:
                plots.append(PlotSpec(name, [ColumnSeries(sheet, x, y, label=y) for y in ys],
                                      "Cycle number", ylabel))
    if _has(book, "conductivity", "file", "conductivity_uS_cm"):
        frame = book.tables["conductivity"]
        if "elapsed_s" in frame and pd.to_numeric(frame["elapsed_s"], errors="coerce").notna().any():
            x, xlabel = "elapsed_s", "Elapsed time (s)"
        elif "source_row" in frame:
            x, xlabel = "source_row", "Source data row (not time)"
        else:
            x = ""
        if x:
            plots.append(PlotSpec("电导率曲线", [ColumnSeries("conductivity", x,
                "conductivity_uS_cm", "file")], xlabel, "Conductivity (μS/cm)"))
    plots.extend(_original_excel_plots(book))
    return plots


def _original_excel_plots(book):
    """Recognize the user's wide Excel layouts without resampling any curve."""
    plots = []
    def wide(sheet, x, name, xlabel, ylabel, legend_suffix="", columns=None):
        frame = book.tables[sheet]
        ys = columns if columns is not None else [c for c in numeric_columns(frame) if c != x]
        if not ys:
            return
        if len(ys) > MAX_SERIES:
            book.notes.append(f"「{sheet}」超过 200 条曲线，请使用尾号分表或按列选择部分曲线。")
            return
        plots.append(PlotSpec(name, [ColumnSeries(sheet, x, y, label=y+legend_suffix) for y in ys], xlabel, ylabel))
    for sheet, frame in book.tables.items():
        first = str(frame.columns[0])
        if first == "Potential" and sheet != "结果栏":
            ylabel = "Current density (A/g)" if Path(book.path).stem.startswith("CV结果") else "Current (Excel input unit)"
            wide(sheet, first, "CV 叠加图", "Potential (V)", ylabel, " mV/s")
        elif sheet == "数据栏" and {"2", "5", "10", "20", "50", "100", "200"}.issubset(frame.columns):
            wide(sheet, first, "CV 叠加图", "Potential (V)", "Current (Excel input unit)", " mV/s")
        elif sheet == "GCD_Data" and first == "Time":
            wide(sheet, first, "GCD 充放电曲线", "Time (s)", "Potential (V)")
        elif sheet == "全量CP数据" or re.fullmatch(r"尾数\d圈", sheet):
            wide(sheet, first, "CP 循环 · " + sheet, "Time (s)", "Potential (V)")
        elif sheet == "原始处理数据" or re.fullmatch(r"尾号\d", sheet):
            wide(sheet, first, "CV 循环 · " + sheet, "Potential (V)", "Current (A)")
        elif sheet in {"电导率", "浓度", "浓度差", "吸附容量", "吸附速率"} and first == "Time":
            ylabel = {"电导率": "Conductivity (μS/cm)", "浓度": "Concentration (mg/L)",
                      "浓度差": "Concentration change (mg/L)", "吸附容量": "Adsorption capacity (mg/g)",
                      "吸附速率": "Average adsorption rate (mg/g/s)"}[sheet]
            wide(sheet, first, sheet + "曲线", "Time (s)", ylabel, " 圈")
        elif sheet == "原始数据" and {"Time", "电导率"}.issubset(frame.columns):
            wide(sheet, "Time", "原始电导率曲线", "Time (s)", "Conductivity (μS/cm)", columns=["电导率"])
        elif sheet == "吸附最大值" and {"圈号", "最大吸附容量"}.issubset(frame.columns):
            plots.append(PlotSpec("各圈最大吸附容量", [ColumnSeries(sheet, "圈号", "最大吸附容量", mode="line_points")],
                                 "Cycle number", "Maximum adsorption capacity (mg/g)"))
        elif sheet == "结果栏":
            for rate in ("2", "5", "10", "20", "50", "100", "200"):
                if {rate, "k1*"+rate}.issubset(frame.columns):
                    plots.append(PlotSpec(rate+" mV/s 电容贡献", [ColumnSeries(sheet, first, rate, label="Total"),
                        ColumnSeries(sheet, first, "k1*"+rate, label="k1 × v")], "Potential (V)", "Current (Excel input unit)"))
        elif sheet == "放电时间汇总" and {"圈号", "放电时长(s)"}.issubset(frame.columns):
            plots.append(PlotSpec("CP 放电时间", [ColumnSeries(sheet, "圈号", "放电时长(s)", mode="line_points")],
                                  "Cycle number", "Discharge time (s; original script)"))
        elif sheet == "结果" and {"比电容", "赝电容"}.issubset(frame.columns):
            plots.append(PlotSpec("电容贡献率", [ColumnSeries(sheet, first, y, label=y, mode="line_points")
                         for y in ("比电容", "赝电容")], "Scan rate (mV/s)", "Contribution (%)"))
    return plots


def generic_plot(sheet, x, ys, group="", mode="line", name="自定义绘图", xlabel="", ylabel=""):
    if not x or not ys:
        raise ValueError("请选择一个 X 列和至少一个 Y 列。")
    return PlotSpec(name, [ColumnSeries(sheet, x, y, group, mode=mode, label=y) for y in ys],
                    xlabel or x, ylabel or (ys[0] if len(ys) == 1 else "Value"))


def _series_data(book, spec):
    series = []
    for item in spec.series:
        if item.sheet not in book.tables:
            raise ValueError(f"缺少工作表：{item.sheet}")
        frame = book.tables[item.sheet]
        required = [item.x, item.y] + ([item.group] if item.group else [])
        required += [item.filter_column] if item.filter_column else []
        if not set(required).issubset(frame.columns):
            raise ValueError(f"工作表「{item.sheet}」缺少绘图列。")
        if item.filter_column:
            frame = frame.loc[frame[item.filter_column].astype(str).str.upper().eq(item.filter_value.upper())]
        if item.group:
            # An entirely blank source row is a line break belonging to the
            # preceding sample, not a new unnamed group.
            blank = frame.isna().all(axis=1)
            if blank.any():
                frame = frame.copy()
                frame.loc[blank, item.group] = frame[item.group].ffill()[blank]
            groups = frame.groupby(item.group, sort=False, dropna=False)
        else:
            groups = [("", frame)]
        for group, part in groups:
            if not len(part):
                continue
            numbers = []
            for column in (item.x, item.y):
                values = part[column]
                converted = pd.to_numeric(values, errors="coerce")
                invalid = values.notna() & values.astype(str).str.strip().ne("") & converted.isna()
                if invalid.any():
                    raise ValueError(f"「{item.sheet}」的「{column}」含非数值（如 {values[invalid].iloc[0]!s}）。请选择数值列。")
                numbers.append(converted.to_numpy(dtype=float, na_value=np.nan))
            x, y = numbers
            if np.isinf(x).any() or np.isinf(y).any():
                raise ValueError("所选列包含无穷大，无法绘图，请核对源数据。")
            if not np.any(np.isfinite(x) & np.isfinite(y)):
                continue
            if item.negate_y:
                y = -y
            label = (item.labels.get(str(group), str(group)) if item.group else item.label or item.y)
            if item.group and item.label:
                label += " · " + item.label
            series.append((x, y, label, item.mode))
            if len(series) > MAX_SERIES:
                raise ValueError("一张图超过 200 条曲线，请按样品拆分表格。")
    if not series:
        raise ValueError("所选列没有成对的有效数值，无法绘图。")
    return series


def draw_plot(book: TableBook, spec: PlotSpec, preset="paper", dpi=100) -> Figure:
    """Create one standalone Figure. All samples remain in their original order."""
    base = PRESETS.get(preset, PRESETS["paper"])
    style = resolve_plot_style(spec.style, preset)
    series = _series_data(book, spec)
    # Preview uses a small rendering DPI; physical dimensions and all appearance
    # settings match export, whose raster DPI is stored in the style.
    figure = Figure(figsize=(style["width_cm"] / 2.54, style["height_cm"] / 2.54),
                    dpi=dpi, layout="constrained", facecolor="white")
    FigureCanvasAgg(figure)
    axes = figure.add_subplot(111)
    color_groups = {}
    handles, labels = [], []
    for index, (x, y, label, mode) in enumerate(series):
        identity = label.removesuffix(" · 拟合")
        color_index = color_groups.setdefault(identity, len(color_groups))
        color = style["palette"][color_index % len(style["palette"])]
        if index < len(style["series_colors"]) and style["series_colors"][index]:
            color = style["series_colors"][index]
        if index < len(style["series_labels"]) and style["series_labels"][index]:
            label = style["series_labels"][index]
        line = ("-", "--", "-.", ":")[index % 4] if base["mono"] else "-"
        marker = ("o", "s", "^", "D", "v")[index % 5]
        if mode == "scatter":
            artist = axes.scatter(x, y, label=label, s=20, marker=marker, facecolors="none", edgecolors=color, linewidths=1)
        elif mode == "bar":
            artist = axes.bar(x, y, label=label, color=color, alpha=0.75)
        else:
            artist, = axes.plot(x, y, label=label, color=color, linestyle=line, linewidth=style["linewidth"],
                                marker=marker if mode == "line_points" else None,
                                markersize=4, markerfacecolor="white")
        handles.append(artist)
        labels.append(label)
    families = list(dict.fromkeys([style["font_family"], *FONT[1:]]))
    label_style = {"fontfamily": families, "fontsize": style["axis_font_size"]}
    axes.set_xlabel(spec.xlabel, **label_style)
    axes.set_ylabel(spec.ylabel, **label_style)
    axes.set_title(spec.title, **label_style)
    axes.tick_params(direction="in", top=True, right=True, labelsize=style["tick_font_size"])
    if style["xlim"] is not None:
        axes.set_xlim(*style["xlim"])
    if style["ylim"] is not None:
        axes.set_ylim(*style["ylim"])
    axes.minorticks_on()
    axes.tick_params(which="minor", direction="in", top=True, right=True)
    for text in axes.get_xticklabels() + axes.get_yticklabels() + [axes.xaxis.get_offset_text(), axes.yaxis.get_offset_text()]:
        text.set_fontfamily(families)
        text.set_fontsize(style["tick_font_size"])
    if base["grid"]:
        axes.grid(alpha=0.2, linewidth=0.5)
    if spec.equal:
        axes.set_aspect("equal", adjustable="box")
    if style["legend_visible"] is True or (style["legend_visible"] is None and len(series) <= 20):
        # Explicit handles also allow labels that begin with an underscore.
        axes.legend(handles, labels, loc=style["legend_location"], frameon=False,
                    prop={"family": families, "size": style["legend_font_size"]})
    return figure


def _safe_name(value):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")[:80]
    return "plot_" + (value or "figure")


def export_plots(book, specs, preset="paper", formats=("png", "pdf", "svg"), output_dir=None) -> str:
    """Export full data serially to a new project output subdirectory + recipe."""
    if not specs:
        raise ValueError("没有可以输出的图。")
    if len(specs) > 64:
        raise ValueError("一次最多输出 64 张图，请分批输出。")
    formats = tuple(formats)
    if not formats or set(formats) - {"png", "pdf", "svg"}:
        raise ValueError("输出格式仅支持 PNG、PDF、SVG。")
    try:
        source_hash = _hash(book.path)
    except OSError as exc:
        raise ValueError("源表格无法读取，请重新导入文件后输出。") from exc
    if source_hash != book.sha256:
        raise ValueError("源表格在导入后已发生变化，请重新导入文件后输出；本次未生成图片。")
    # Validate all plots before making any output.
    rendered_specs = []
    for spec in specs:
        _series_data(book, spec)
        rendered_specs.append(replace(spec, style=resolve_plot_style(spec.style, preset)))
    from ec_export import validate_output_dir, unused_path
    if output_dir is not None:
        from ec_output_dialog import validate_names
        validate_names([(str(i), spec.name) for i, spec in enumerate(specs)])
        output = Path(validate_output_dir(output_dir, [str(Path(book.path).parent)]))
    else:
        output = BASE_DIR / "output" / ("绘图_" + datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + uuid.uuid4().hex[:8])
    output.mkdir(parents=True, exist_ok=True)
    records = []
    try:
        for index, spec in enumerate(rendered_specs, 1):
            figure = draw_plot(book, spec, preset, dpi=100)
            files = []
            try:
                for extension in formats:
                    filename = (spec.name if output_dir is not None else f"{index:02d}_{_safe_name(spec.name)}") + "." + extension
                    target = Path(unused_path(output / filename))
                    filename = target.name
                    figure.savefig(target, dpi=spec.style["dpi"], facecolor="white")
                    files.append(filename)
            finally:
                figure.clear()
            records.append({"spec": asdict(spec), "files": files})
        dpis = {spec.style["dpi"] for spec in rendered_specs}
        recipe = {"schema": "ec-lab-plot-recipe", "version": 1,
                  "source": {"path": book.path, "sha256": book.sha256},
                  "preset": preset, "preset_values": PRESETS.get(preset, PRESETS["paper"]),
                  "dpi": next(iter(dpis)) if len(dpis) == 1 else None,
                  "dpi_policy": "each plot's final style stores its export DPI",
                  "formats": list(formats), "plots": records,
                  "row_order": "source order; no sorting, smoothing, resampling or truncation",
                  "missing_values": "blank cells remain NaN gaps; nonnumeric values are rejected",
                  "legend": "per-plot style; null visibility shows legends for at most 20 series",
                  "style_attribution": STYLE_SOURCE}
        recipe["software"] = {"python": sys.version.split()[0], "matplotlib": matplotlib.__version__,
                              "numpy": np.__version__, "pandas": pd.__version__}
        Path(unused_path(output / "绘图配方.json")).write_text(json.dumps(recipe, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as exc:
        (output / "未完成.txt").write_text("输出未完成，请勿把此目录当作完整结果。\n" + str(exc), encoding="utf-8")
        raise RuntimeError(f"输出未完成：{exc}\n已生成部分文件的位置：{output}") from exc
    return str(output)


def load_recipe(path):
    """Load declarative JSON only. Recheck source hash before reproducible export."""
    recipe_path = Path(path)
    if recipe_path.stat().st_size > 2 * 1024 * 1024:
        raise ValueError("绘图配方超过 2 MB，无法导入。")
    recipe = json.loads(recipe_path.read_text(encoding="utf-8"))
    if recipe.get("schema") != "ec-lab-plot-recipe" or recipe.get("version") != 1:
        raise ValueError("这不是受支持的 ec-lab 绘图配方。")
    source = recipe["source"]
    book = load_table_book(source["path"])
    if book.sha256 != source["sha256"]:
        raise ValueError("原始表格已发生变化，无法按旧配方复现。请重新导入表格。")
    _normalize_legacy_summary(book)
    specs = []
    for record in recipe["plots"]:
        values = dict(record["spec"])
        values["style"] = validate_plot_style(values.get("style", {}))
        values["series"] = [ColumnSeries(**value) for value in values["series"]]
        for series in values["series"]:
            if (series.sheet == "summary" and not series.group and
                    series.filter_column == "type" and series.filter_value.upper() in ("CV", "GCD") and
                    series.mode == "line_points" and
                    series.x in ("scan_rate_mVs", "current_density_A_g") and
                    _has(book, "summary", "type", series.x)):
                series.mode = _summary_mode(book, series.filter_value.upper(), series.x)
        specs.append(PlotSpec(**values))
    if recipe.get("preset") not in PRESETS:
        raise ValueError("配方的绘图模板不可用。")
    saved_style = recipe.get("preset_values")
    current_style = json.loads(json.dumps(PRESETS[recipe["preset"]], ensure_ascii=False))
    if saved_style is not None and saved_style != current_style:
        raise ValueError("绘图模板的定义已经变化，无法精确复现旧配方。请重新选择模板后输出。")
    return book, specs, recipe["preset"]
