"""Fill the supplied fixed-grid contribution template without restyling it."""
from pathlib import Path
import hashlib
import re
import zipfile

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as letter

from ec_legacy import TEMPLATE, _cache_formulas
from ec_export import unused_path

RATES = np.array([2., 5., 10., 20., 50., 100., 200.])
GRID = np.concatenate((np.linspace(-.8, .8, 1601), np.linspace(.799, -.8, 1600)))


def template_caches(x, currents, label="Potential"):
    """Numeric cache for the unchanged POWER/SLOPE/INTERCEPT/SUM formulas.

    The regression uses the exact scan-rate units and signed trapezoidal
    integration in the supplied workbook, without clipping percentages.
    """
    roots = np.sqrt(RATES)
    normalized = currents/roots
    centered = roots-roots.mean()
    slopes = normalized @ centered / np.sum(centered**2)
    intercepts = normalized.mean(axis=1)-slopes*roots.mean()
    capacitive = slopes[:, None]*RATES
    integrals = (currents[:-1]+currents[1:])/2*np.diff(x)[:, None]
    cap_integrals = (capacitive[:-1]+capacitive[1:])/2*np.diff(x)[:, None]
    total, cap = integrals.sum(axis=0), cap_integrals.sum(axis=0)
    if np.any(np.isclose(total, 0, atol=1e-15)):
        raise ValueError("CV 有符号积分为零，原模板的贡献率公式会除零。请核对电流与扫描顺序。")
    processing, result, summary = {}, {"A1": label}, {}
    for i, rate in enumerate(RATES):
        for col in (10+i, 18+i, 28+i):
            processing[f"{letter(col)}2"] = rate
        processing[f"{letter(10+i)}3"] = roots[i]
    for index in range(len(x)):
        row = index+4
        processing[f"A{row}"] = x[index]
        for j in range(7):
            processing[f"{letter(j+2)}{row}"] = currents[index,j]
        row = index+3
        processing[f"Z{row}"], processing[f"AA{row}"] = slopes[index], intercepts[index]
        for j in range(7):
            processing[f"{letter(j+18)}{row}"] = normalized[index,j]
            processing[f"{letter(j+28)}{row}"] = capacitive[index,j]
        row = index+2
        result[f"A{row}"] = x[index]
        for j in range(7):
            result[f"{letter(j+2)}{row}"] = currents[index,j]
            result[f"{letter(j+9)}{row}"] = capacitive[index,j]
            if index < len(x)-1:
                result[f"{letter(j+17)}{row}"] = integrals[index,j]
                result[f"{letter(j+24)}{row}"] = cap_integrals[index,j]
    # The template contains one extra regression row referencing blank cells.
    for col in [*range(18,25), 26,27, *range(28,35)]:
        processing[f"{letter(col)}3204"] = 0.
    for index in range(7):
        percentage = cap[index]/total[index]*100
        for col, value in (("B", total[index]), ("C", cap[index]), ("D", percentage), ("E", 100-percentage)):
            summary[f"{col}{index+2}"] = value
    return {"xl/worksheets/sheet2.xml": processing,
            "xl/worksheets/sheet3.xml": result, "xl/worksheets/sheet4.xml": summary}


def fill_template(source, folder):
    source, folder = Path(source), Path(folder)
    if hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() != "7fad959cfb49ddefe2cab7b71f39ad9504d6501382e5911d2599b8098e2729f3":
        raise ValueError("内置贡献率模板已被修改，请恢复原模板后再生成，避免公式与预计算结果不一致。")
    if source.suffix.lower() != ".xlsx":
        raise ValueError("贡献率模板需要 CV结果.xlsx。")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["数据栏"] if "数据栏" in workbook.sheetnames else workbook.worksheets[0]
        if sheet.max_row != 3202 or sheet.max_column != 8:
            raise ValueError("原模板需要 3201 行数据和 8 列（电压＋七种扫速），本文件尺寸不符。")
        rows = list(sheet.values)
    finally:
        workbook.close()
    try:
        headers = np.asarray(rows[0][1:], dtype=float)
        values = np.asarray(rows[1:], dtype=float)
    except (TypeError, ValueError) as exc:
        raise ValueError("电压、电流与扫速必须是数值。") from exc
    if not np.isfinite(values).all() or set(headers) != set(RATES):
        raise ValueError("存在缺失值或扫速不符合原模板的 2/5/10/20/50/100/200 mV/s。")
    x = values[:,0]
    if not np.allclose(x, GRID, atol=1e-7, rtol=0):
        raise ValueError("电压序列与原模板的 −0.8→0.8→−0.8 V、0.001 V 步长不一致；为避免改变数据，不自动插值或倒序。")
    currents = values[:, [list(headers).index(rate)+1 for rate in RATES]]
    caches = template_caches(x, currents)
    destination = Path(unused_path(folder / f"{source.stem}_电容贡献率.xlsx"))
    with zipfile.ZipFile(TEMPLATE) as template, zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as output:
        for entry in template.infolist():
            raw = template.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                def fill(match):
                    opening, body = match.group(1), match.group(2)
                    coordinate = re.search(r'\br="([^"]+)"', opening).group(1)
                    cell = re.fullmatch(r"([A-H])(\d+)", coordinate)
                    if not cell or int(cell[2]) < 2:
                        return match.group(0)
                    col, row = ord(cell[1])-65, int(cell[2])-2
                    value = x[row] if col==0 else currents[row,col-1]
                    body = re.sub(r"<v(?:\s[^>]*)?>.*?</v>|<v\s*/>", "", body, flags=re.S)
                    opening = re.sub(r'\s+t="[^"]*"', "", opening)
                    return opening+body+f"<v>{float(value)!r}</v></c>"
                raw = re.sub(r"(<c\b[^>]*>)(.*?)</c>", fill, raw.decode("utf-8"), flags=re.S).encode("utf-8")
            elif entry.filename == "xl/worksheets/sheet4.xml":
                # The original result sheet used default narrow columns, so
                # its full-precision cached numbers overlapped. Keep fonts,
                # formats and formula positions, only make those columns fit.
                text = raw.decode("utf-8")
                columns = ('<cols><col min="1" max="1" width="10" customWidth="1"/>'
                           '<col min="2" max="3" width="24" customWidth="1"/>'
                           '<col min="4" max="5" width="20" customWidth="1"/></cols>')
                text = re.sub(r"<cols>.*?</cols>", "", text, flags=re.S)
                raw = text.replace("<sheetData>", columns+"<sheetData>", 1).encode("utf-8")
            output.writestr(entry, raw)
    _cache_formulas(destination, caches)
    return destination
