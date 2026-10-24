# -*- coding: utf-8 -*-
"""Focused tests for independent Excel plotting (synthetic data only)."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import numpy as np
import pandas as pd
from openpyxl import Workbook
from PIL import Image

import ec_plotting as plot

TEST_ROOT = Path(__file__).resolve().parents[1] / 'output' / '_test_tmp'
TEST_ROOT.mkdir(parents=True, exist_ok=True)


class PlottingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=TEST_ROOT)
        self.root = Path(self.temp.name)
        self.path = self.root / "processed_data.xlsx"
        with pd.ExcelWriter(self.path, engine="openpyxl") as writer:
            pd.DataFrame({"file": ["CV_a"] * 3, "Potential_V": [0, 1, 0],
                          "CurrentDensity_A_g": [1, 2, -1]}).to_excel(writer, sheet_name="CV", index=False)
            pd.DataFrame({"file": ["GCD_a"] * 3, "time_s": [0, 2, 5],
                          "Potential_V": [0, 1, 0]}).to_excel(writer, sheet_name="GCD", index=False)
            pd.DataFrame({"type": ["CV", "GCD"], "file": ["CV_a", "GCD_a"],
                          "condition": ["10 mV/s", "2 A/g"], "capacitance_F_g": [100, 90]}).to_excel(writer, sheet_name="summary", index=False)
            pd.DataFrame({"file": ["EIS_a"] * 3, "Freq_Hz": [100, 10, 1],
                          "Z_re_ohm": [1, 2, 3], "Z_im_ohm": [-1, -2, -1]}).to_excel(writer, sheet_name="EIS", index=False)
            pd.DataFrame({"file": ["EIS_a"] * 2, "Freq_Hz": [50, 5],
                          "Z_re_fit_ohm": [1.5, 2.5], "Z_im_fit_ohm": [-1.2, -1.8]}).to_excel(writer, sheet_name="EIS_fit_curve", index=False)
            pd.DataFrame({"cycle": [10, 20], "sample": [100, 95]}).to_excel(writer, sheet_name="cycle_capacitance", index=False)
            pd.DataFrame({"cycle": [10, 20], "sample": [100, 95]}).to_excel(writer, sheet_name="cycle_retention", index=False)
            pd.DataFrame({"file": ["cond_a"] * 3, "elapsed_s": [0, 1, 2],
                          "conductivity_uS_cm": [100, None, 120], "source_row": [2, 3, 4]}).to_excel(writer, sheet_name="conductivity", index=False)
        self.book = plot.load_table_book(str(self.path))

    def tearDown(self):
        self.temp.cleanup()

    def test_auto_recognizes_old_and_new_tables_without_recalculation(self):
        specs = plot.automatic_plots(self.book)
        self.assertEqual(len(specs), 8)
        cv = plot.draw_plot(self.book, specs[0])
        np.testing.assert_array_equal(cv.axes[0].lines[0].get_xdata(), [0, 1, 0])
        self.assertEqual(cv.axes[0].lines[0].get_label(), "10 mV/s")
        cv.clear()
        eis = next(spec for spec in specs if spec.name.startswith("EIS"))
        figure = plot.draw_plot(self.book, eis)
        self.assertEqual(len(figure.axes[0].collections[0].get_offsets()), 3)
        np.testing.assert_array_equal(figure.axes[0].lines[0].get_ydata(), [1.2, 1.8])
        figure.clear()

    def test_grouped_order_and_blank_line_gap(self):
        source = self.root / "blank.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["sample", "x", "y"])
        sheet.append(["b", 3, 10])
        sheet.append([None, None, None])
        sheet.append(["b", 1, 20])
        sheet.append(["a", 2, 30])
        workbook.save(source)
        book = plot.load_table_book(str(source))
        spec = plot.generic_plot("Sheet", "x", ["y"], "sample")
        series = plot._series_data(book, spec)
        self.assertEqual([item[2] for item in series], ["b · y", "a · y"])
        np.testing.assert_array_equal(series[0][0], [3, np.nan, 1])
        self.assertTrue(np.isnan(series[0][1][1]))

    def test_generic_multiple_y_and_invalid_data(self):
        source = self.root / "data.csv"
        source.write_text("x,y,z\n3,1,4\n1,2,5\n2,3,6\n", encoding="utf-8")
        book = plot.load_table_book(str(source))
        spec = plot.generic_plot("data", "x", ["y", "z"])
        series = plot._series_data(book, spec)
        self.assertEqual(len(series), 2)
        np.testing.assert_array_equal(series[0][0], [3, 1, 2])
        book.tables["data"]["y"] = [1, "oops", 2]
        with self.assertRaisesRegex(ValueError, "非数值"):
            plot.draw_plot(book, spec)

    def test_export_full_resolution_recipe_and_no_input_overwrite(self):
        source_hash = plot._hash(self.path)
        spec = plot.automatic_plots(self.book)[0]
        with patch.object(plot, "BASE_DIR", self.root):
            output = Path(plot.export_plots(self.book, [spec]))
            again = Path(plot.export_plots(self.book, [spec], formats=("svg",)))
        self.assertNotEqual(output, again)
        self.assertEqual(plot._hash(self.path), source_hash)
        self.assertEqual(len(list(output.glob("*.png"))), 1)
        self.assertEqual(len(list(output.glob("*.pdf"))), 1)
        self.assertEqual(len(list(output.glob("*.svg"))), 1)
        with Image.open(next(output.glob("*.png"))) as image:
            self.assertEqual(image.size, (3840, 2700))
        book, specs, preset = plot.load_recipe(output / "绘图配方.json")
        self.assertEqual(book.sha256, source_hash)
        self.assertEqual(specs[0].series[0].x, "Potential_V")
        self.assertEqual(preset, "paper")
        # A changed source is explicitly rejected during reproduction.
        with open(self.path, "ab") as handle:
            handle.write(b"changed")
        with self.assertRaisesRegex(ValueError, "发生变化"):
            plot.load_recipe(output / "绘图配方.json")

    def test_limit_is_rejection_not_truncation(self):
        with patch.object(plot, "MAX_ROWS", 2):
            with self.assertRaises(ValueError):
                plot.load_table_book(str(self.path))

    def test_legacy_summary_capacitance_recipe_replays_derived_columns(self):
        self.assertNotIn("scan_rate_mVs", self.book.tables["summary"])
        specs = [spec for spec in plot.automatic_plots(self.book)
                 if spec.name in ("CV 比电容图", "GCD 比电容图")]
        with patch.object(plot, "BASE_DIR", self.root):
            first = Path(plot.export_plots(self.book, specs, formats=("svg",)))
            restored, restored_specs, preset = plot.load_recipe(first / "绘图配方.json")
            second = Path(plot.export_plots(restored, restored_specs, preset, formats=("svg",)))
        self.assertEqual(len(list(second.glob("*.svg"))), 2)
        np.testing.assert_array_equal(plot._series_data(restored, restored_specs[0])[0][0], [10])
        np.testing.assert_array_equal(plot._series_data(restored, restored_specs[1])[0][0], [2])

    def test_repeated_conditions_are_scatter_and_explained(self):
        repeated = pd.DataFrame({"type": ["CV", "GCD"], "file": ["another_cv", "another_gcd"],
                                 "condition": ["10 mV/s", "2 A/g"], "capacitance_F_g": [120, 110]})
        self.book.tables["summary"] = pd.concat([self.book.tables["summary"], repeated], ignore_index=True)
        specs = [spec for spec in plot.automatic_plots(self.book)
                 if spec.name in ("CV 比电容图", "GCD 比电容图")]
        self.assertEqual([spec.series[0].mode for spec in specs], ["scatter", "scatter"])
        self.assertTrue(any("CV 存在重复" in note and "散点" in note for note in self.book.notes))
        self.assertTrue(any("GCD 存在重复" in note and "散点" in note for note in self.book.notes))
        for spec in specs:
            figure = plot.draw_plot(self.book, spec)
            self.assertEqual(len(figure.axes[0].lines), 0)
            self.assertEqual(len(figure.axes[0].collections[0].get_offsets()), 2)
            figure.clear()

    def test_export_rejects_changed_source_before_creating_output(self):
        specs = plot.automatic_plots(self.book)
        with open(self.path, "ab") as handle:
            handle.write(b"source changed after import")
        with patch.object(plot, "BASE_DIR", self.root):
            with self.assertRaisesRegex(ValueError, "重新导入"):
                plot.export_plots(self.book, specs, formats=("svg",))
        self.assertFalse((self.root / "output").exists())


class StudioTests(unittest.TestCase):
    def test_background_import_and_single_canvas(self):
        from PySide6.QtCore import QEventLoop, QTimer
        from PySide6.QtWidgets import QApplication
        from ec_plot_studio import PlotStudio
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory(dir=TEST_ROOT) as directory:
            source = Path(directory) / "generic.csv"
            source.write_text("x,y\n0,1\n1,2\n", encoding="utf-8")
            studio = PlotStudio()
            studio.load_path(str(source))
            self.assertTrue(studio._busy)
            loop = QEventLoop()
            studio.worker.finished.connect(loop.quit)
            QTimer.singleShot(15000, loop.quit)
            loop.exec()
            self.assertFalse(studio._busy)
            self.assertIsNotNone(studio.book)
            self.assertEqual(studio.tabs.currentIndex(), 1)
            studio._preview()
            first = studio.canvas
            studio._preview()
            self.assertIsNot(studio.canvas, first)
            self.assertEqual(len(studio.canvas.figure.axes), 1)
            studio.reject()
            app.processEvents()


if __name__ == "__main__":
    unittest.main(verbosity=2)
