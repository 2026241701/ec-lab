"""Compatibility tests against the actual, unchanged user scripts."""
from contextlib import redirect_stdout
import io
import copy
import hashlib
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("OPENBLAS_NUM_THREADS", "1")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
import numpy as np
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QMessageBox
from PySide6.QtCore import Qt

from ec_legacy import electrochemistry, cp_cycles, cv_cycles, cdi_workbook, ec, cp, cv, run_legacy
from ec_contribution import fill_template, GRID, RATES
from ec_plotting import load_table_book, automatic_plots, _series_data
from ec_output_dialog import OutputSelectionDialog, render_selected, offer_open_folder, validate_names
import make_mock
import eclab

ROOT = Path(__file__).resolve().parents[1]


class LegacyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=ROOT / "output")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.inputs = self.root / "source"
        self.inputs.mkdir()
        self.expected, self.actual = self.root / "expected", self.root / "actual"
        self.expected.mkdir(); self.actual.mkdir()
        wave = np.tile(np.r_[np.linspace(-.8, .8, 161)[:-1], np.linspace(.8, -.8, 161)[:-1]], 4)
        self.cv_paths = []
        for rate in (2, 5):
            path = self.inputs / f"CV_{rate}.txt"
            path.write_text("Potential/V, Current/A\n" + "\n".join(f"{v:.6f}, {rate*.0001+v*.00001:.9f}" for v in wave))
            self.cv_paths.append(path)
        self.cp_path = self.inputs / "GCD_1.txt"
        self.cp_path.write_text("Time/sec, Potential/V\n" + "\n".join(f"{i*.1:.2f}, {v:.6f}" for i,v in enumerate(wave)) + "\nHold Time/sec, Current/A\n")

    def compare_workbooks(self, expected, actual):
        a, b = load_workbook(expected), load_workbook(actual)
        try:
            self.assertEqual(a.sheetnames, b.sheetnames)
            for sheet in a:
                self.assertEqual(list(sheet.values), list(b[sheet.title].values), sheet.title)
                for row in sheet:
                    for cell in row:
                        other = b[sheet.title][cell.coordinate]
                        self.assertEqual(copy.copy(cell.fill), copy.copy(other.fill))
                        self.assertEqual(copy.copy(cell.font), copy.copy(other.font))
                        self.assertEqual(cell.number_format, other.number_format)
        finally:
            a.close(); b.close()

    def test_electrochemistry_matches_original_cells_and_styles(self):
        with redirect_stdout(io.StringIO()), patch.object(ec, "Path", return_value=self.inputs), \
             patch.object(ec, "OUTPUT_GCD", str(self.expected / "GCD结果.xlsx")), \
             patch.object(ec, "OUTPUT_CV", str(self.expected / "CV结果.xlsx")), patch("builtins.input", return_value="2"):
            ec.process_gcd(); ec.process_cv()
        electrochemistry(self.cv_paths+[self.cp_path], self.actual, 2, progress=lambda _:None)
        for name in ("CV结果.xlsx", "GCD结果.xlsx"):
            self.compare_workbooks(self.expected/name, self.actual/name)
            book = load_table_book(str(self.actual/name))
            self.assertEqual(len(automatic_plots(book)), 1)

    def test_cycle_workbooks_match_original(self):
        with redirect_stdout(io.StringIO()):
            v = cp.process_raw_txt(str(self.cp_path)); troughs = cp.find_all_troughs(v)
            cycles = cp.split_cycles(v, troughs)
            cp.export_xlsx(cp.build_export_df(cycles), cp.calc_discharge_time(cycles,v,troughs), str(self.expected/'cp.xlsx'))
            actual = cp_cycles(self.cp_path, self.actual)
            self.compare_workbooks(self.expected/'cp.xlsx', actual)
            v,c = cv.read_and_clean_txt(str(self.cv_paths[0]))
            cycles = cv.split_cycles(v,c,cv.detect_points(v,'peak'))
            cv.save_xlsx_all(cv.filter_by_last_digit(cv.build_processed_table(cycles)), str(self.expected/'cv'))
            actual = cv_cycles(self.cv_paths[0], self.actual, progress=lambda _:None)
            self.compare_workbooks(self.expected/'cv.xlsx', actual)
            for spec in automatic_plots(load_table_book(str(actual))):
                self.assertTrue(_series_data(load_table_book(str(actual)), spec))

    def test_cdi_keeps_formulas_and_caches_but_blanks_missing_and_zero_time(self):
        path = self.inputs/'conductivity.csv'
        path.write_text('time_s,conductivity_uS_cm\n0,100\n5,90\n10,80\n15,70\n20,60\n')
        result = cdi_workbook(path,self.actual,{'ads_minutes':1/6,'des_minutes':1/6},progress=lambda _:None)
        formula, cached = load_workbook(result), load_workbook(result,data_only=True)
        self.addCleanup(formula.close);self.addCleanup(cached.close)
        self.assertEqual(formula.sheetnames,['原始数据','电导率','浓度','浓度差','吸附容量','吸附速率','吸附最大值'])
        self.assertEqual(formula['吸附容量']['B3'].value,'=浓度差!B3*0.05/(原始数据!$D$1)*1000')
        expected = 10/1.97252*0.05/24*1000
        self.assertAlmostEqual(cached['吸附容量']['B3'].value,expected)
        self.assertAlmostEqual(cached['吸附速率']['B3'].value,expected/5)
        self.assertIn(cached['吸附速率']['B2'].value,(None,''))
        for sheet in ('浓度','浓度差','吸附容量','吸附速率'):
            self.assertIn(cached[sheet]['C3'].value,(None,''))
        book=load_table_book(str(result))
        self.assertEqual(len(automatic_plots(book)),7)

    def test_template_formulas_styles_and_calculation(self):
        from openpyxl import Workbook
        import zipfile
        from ec_legacy import TEMPLATE
        wb=Workbook();ws=wb.active
        ws.append(['Potential',*RATES])
        branch=np.r_[np.ones(1601),-np.ones(1600)]
        currents=branch[:,None]*(RATES*.01+np.sqrt(RATES)*.03)
        for x, row in zip(GRID,currents):ws.append([float(x),*map(float,row)])
        source=self.inputs/'CV结果.xlsx';wb.save(source)
        output=fill_template(source,self.actual)
        with zipfile.ZipFile(TEMPLATE) as original,zipfile.ZipFile(output) as exported:
            self.assertEqual(original.read('xl/styles.xml'),exported.read('xl/styles.xml'))
            import re
            for i in (2,3,4):
                name=f'xl/worksheets/sheet{i}.xml'
                self.assertEqual(hashlib.sha256(b''.join(re.findall(rb'<f\b[^>]*?/>|<f\b[^>]*>.*?</f>',original.read(name)))).hexdigest(), hashlib.sha256(b''.join(re.findall(rb'<f\b[^>]*?/>|<f\b[^>]*>.*?</f>',exported.read(name)))).hexdigest())
        result=load_workbook(output,data_only=True,read_only=True)
        for index,rate in enumerate(RATES,2):
            self.assertAlmostEqual(result['结果'].cell(index,4).value,rate*.01/(rate*.01+np.sqrt(rate)*.03)*100,places=8)
        result.close()
        self.assertEqual(len(automatic_plots(load_table_book(str(output)))),9)

    def test_source_directory_protection_and_duplicate_conditions(self):
        with self.assertRaises(ValueError):
            run_legacy('electrochem',self.cv_paths,self.inputs,{'mass_mg':2})
        other=self.inputs/'extra_2.txt';other.write_bytes(self.cv_paths[0].read_bytes())
        with self.assertRaises(ValueError):
            electrochemistry([self.cv_paths[0],other],self.actual,2)


class OutputDialogTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_selection_validation_and_cancel(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'output') as folder:
            dlg=OutputSelectionDialog(None,[('a','CV 叠加图'),('b','GCD 充放电曲线')],folder)
            dlg.select_all(False)
            self.assertFalse(dlg.save_button.isEnabled())
            dlg.table.item(1,0).setCheckState(Qt.Checked)
            dlg.table.item(1,1).setText('我的充放电图')
            dlg.accept()
            self.assertEqual(dlg.selection,[('b','我的充放电图')])
            self.assertEqual(list(Path(folder).iterdir()),[])
            dlg.deleteLater()
            cancelled=OutputSelectionDialog(None,[('a','CV')],folder)
            cancelled.reject()
            self.assertEqual(list(Path(folder).iterdir()),[])
            cancelled.deleteLater()
        for names in ([], [('a','CON')], [('a','../escape')], [('a','same'),('b','SAME')]):
            with self.assertRaises(ValueError):validate_names(names)

    def test_only_selected_output_and_folder_prompt(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'output') as folder:
            cfg=make_mock.sample_cfg({'eis':{'fit':False},'output':{'dir':folder,'formats':['png'],'dpi':72}})
            with redirect_stdout(io.StringIO()):data=eclab.prepare(cfg)
            files=render_selected(data,cfg,[('CV_overlay','我的图')])
            self.assertEqual([Path(p).name for p in files],['我的图.png'])
            self.assertEqual(len(list(Path(folder).iterdir())),1)
            render_selected(data,cfg,[('CV_overlay','我的图')])
            self.assertEqual(len(list(Path(folder).iterdir())),2)
            with patch.object(QMessageBox,'question',return_value=QMessageBox.No),patch('os.startfile') as open_folder:
                offer_open_folder(None,folder,1);open_folder.assert_not_called()
            with patch.object(QMessageBox,'question',return_value=QMessageBox.Yes),patch('os.startfile') as open_folder:
                offer_open_folder(None,folder,1);open_folder.assert_called_once_with(folder)


if __name__=='__main__':unittest.main()
