"""Small deterministic import/export guards. Never writes to E:\\Data."""
import hashlib
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook, load_workbook
from ec_conductivity import export_conductivity, process_files, read_conductivity, validate_destination


class ConductivityTests(unittest.TestCase):
    def setUp(self):
        scratch = Path(__file__).resolve().parents[1] / "output" / "_test_temp"
        scratch.mkdir(parents=True, exist_ok=True)
        self.tmp = tempfile.TemporaryDirectory(prefix="ec_cond_test_", dir=scratch)
        self.root = Path(self.tmp.name)
        self.source = self.root / "source"
        self.source.mkdir()

    def tearDown(self):
        self.tmp.cleanup()

    def text_file(self, text, name="data.csv"):
        path = self.source / name
        path.write_text(text, encoding="utf-8")
        return path

    def test_unit_conversion_and_originals(self):
        path = self.text_file("测量时间,检测结果\n2026-07-04 21:10:53,1002uS/cm\n2026-07-04 21:10:58,1.2mS/cm\n2026-07-04 21:11:03,0.1S/m\n")
        result = read_conductivity(path)
        self.assertEqual(result.data.elapsed_s.tolist(), [0, 5, 10])
        self.assertEqual(result.data.conductivity_uS_cm.tolist(), [1002, 1200, 1000])
        self.assertEqual(result.data.source_row.tolist(), [2, 3, 4])
        self.assertEqual(result.data.raw_conductivity.tolist(), ["1002uS/cm", "1.2mS/cm", "0.1S/m"])
        self.assertEqual(result.valid_count, 3)

    def test_headerless_requires_unit_and_keeps_zero(self):
        path = self.text_file("0\t0\n5\t100\n", "data.txt")
        missing = read_conductivity(path)
        self.assertEqual(len(missing.data), 2)
        self.assertEqual(missing.valid_count, 0)
        result = read_conductivity(path, "uS/cm")
        self.assertEqual(result.valid_count, 2)
        self.assertEqual(result.data.elapsed_s.tolist(), [0, 5])

    def test_bad_rows_remain_and_time_not_sorted(self):
        path = self.text_file("time,conductivity\n0,10uS/cm\n5,bad\n2,20uS/cm\n,30uS/cm\n10,-2uS/cm\n")
        result = read_conductivity(path)
        self.assertEqual(len(result.data), 5)
        self.assertEqual(result.data.elapsed_s.iloc[:3].tolist(), [0, 5, 2])
        self.assertEqual(result.valid_count, 2)
        self.assertTrue(any("倒退" in issue["message"] for issue in result.issues))

    def test_blank_measurement_rows_are_reported_and_preserved(self):
        path = self.text_file("time,conductivity\n0,10uS/cm\n,\n5,9uS/cm\n")
        result = read_conductivity(path)
        self.assertEqual(len(result.data), 3)
        self.assertEqual(result.valid_count, 2)
        self.assertEqual(result.data.source_row.tolist(), [2, 3, 4])
        row_issues = [i for i in result.issues if i["source_row"] == 3]
        self.assertEqual(len(row_issues), 2)

    def test_explicit_header_and_cell_units_take_precedence(self):
        path = self.text_file("time (min),conductivity (mS/cm)\n1,1\n2,0.1S/m\n")
        result = read_conductivity(path, "uS/cm")
        self.assertEqual(result.data.elapsed_s.tolist(), [0, 60])
        self.assertEqual(result.data.conductivity_uS_cm.tolist(), [1000, 1000])

    def test_duplicate_column_is_reported(self):
        path = self.text_file("time,conductivity,conductivity\n0,1uS/cm,2uS/cm\n")
        result = read_conductivity(path)
        self.assertTrue(result.data.empty)
        self.assertTrue(any("多个" in issue["message"] for issue in result.issues))

    def test_unknown_unit_stays_invalid_even_with_default(self):
        path = self.text_file("time,conductivity\n0,1mg/L\n")
        self.assertEqual(read_conductivity(path, "uS/cm").valid_count, 0)

    def test_xlsx_reordered_columns(self):
        path = self.source / "instrument.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["检测结果", "样品编号", "测量时间"])
        sheet.append(["900μS/cm", "A", "2026-07-04 21:10:53"])
        sheet.append(["850µS/cm", "A", "2026-07-04 21:10:58"])
        book.save(path)
        book.close()
        result = read_conductivity(path)
        self.assertEqual(result.data.elapsed_s.tolist(), [0, 5])
        self.assertEqual(result.data.conductivity_uS_cm.tolist(), [900, 850])

    def test_output_protection_and_no_overwrite(self):
        path = self.text_file("time,conductivity\n0,1uS/cm\n")
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        result = read_conductivity(path)
        with self.assertRaises(ValueError):
            export_conductivity(result, self.source / "result.xlsx")
        with self.assertRaises(ValueError):
            export_conductivity(result, self.source / "nested" / "result.xlsx")
        target = self.root / "output" / "result.xlsx"
        export_conductivity(result, target)
        with self.assertRaises(ValueError):
            export_conductivity(result, target)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)
        book = load_workbook(target, read_only=True)
        self.assertEqual(book.sheetnames, ["conductivity", "issues", "sources", "说明"])
        self.assertEqual(book["conductivity"].max_row, 2)
        book.close()

    def test_raw_formula_like_text_not_executed(self):
        path = self.text_file('time,conductivity,备注\n0,1uS/cm,=1+1\n')
        result = read_conductivity(path)
        target = export_conductivity(result, self.root / "out" / "safe.xlsx")
        book = load_workbook(target, data_only=False)
        sheet = book["conductivity"]
        headers = [c.value for c in sheet[1]]
        cell = sheet.cell(2, headers.index("raw_3_备注") + 1)
        self.assertEqual(cell.value, "=1+1")
        self.assertEqual(cell.data_type, "s")
        book.close()

    def test_batch_continues_after_missing_file(self):
        path = self.text_file("time,conductivity\n0,1uS/cm\n")
        result = process_files([self.source / "missing.xlsx", path])
        self.assertEqual(result.valid_count, 1)
        self.assertTrue(any("文件未导入" in issue["message"] for issue in result.issues))


if __name__ == "__main__":
    unittest.main()
