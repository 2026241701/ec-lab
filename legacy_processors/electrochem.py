
# -*- coding: utf-8 -*-
"""
GCD/CV Processor V2 - 自动识别 + CV归一化
Python 3.10+
"""

import re
from pathlib import Path
import numpy as np
import pandas as pd
from scipy.signal import find_peaks
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# ===== 用户配置 =====
GCD_START_CYCLE = 2
GCD_END_CYCLE = 3
CV_TOLERANCE = 0.001
GCD_TIME_STEP = 0.1
PEAK_MIN_DISTANCE = 20
PEAK_PROMINENCE = 0.05
OUTPUT_GCD = "GCD结果.xlsx"
OUTPUT_CV = "CV结果.xlsx"
# ===================

def read_text(path):
    for enc in ("utf-8","gb18030","gbk","latin1"):
        try:
            return path.read_text(encoding=enc, errors="ignore")
        except:
            pass
    return path.read_text(errors="ignore")

def detect_type(text):
    if "Hold Time" in text:
        return "GCD"
    if "Potential/V, Current/A" in text:
        return "CV"
    return "UNKNOWN"

def simplify_name(path):
    m = re.search(r"(\d+)", path.stem)
    return int(m.group(1)) if m else path.stem

def extract_gcd(path):
    lines = read_text(path).splitlines()
    start,end=None,None
    for i,l in enumerate(lines):
        if "Time/sec, Potential/V" in l:
            start=i+1
        if "Hold Time/sec, Current/A" in l:
            end=i
            break
    if start is None:
        raise ValueError(f"{path.name} 未找到GCD数据")
    data=[]
    for l in lines[start:end]:
        p=[x.strip() for x in l.split(",")]
        if len(p)<2: continue
        try:
            data.append((float(p[0]),float(p[1])))
        except:
            pass
    return pd.DataFrame(data,columns=["Time","Voltage"])

def extract_cv(path):
    lines = read_text(path).splitlines()
    start=None
    for i,l in enumerate(lines):
        if "Potential/V, Current/A" in l:
            start=i+1
            break
    if start is None:
        raise ValueError(f"{path.name} 未找到CV数据")
    data=[]
    for l in lines[start:]:
        p=[x.strip() for x in l.split(",")]
        if len(p)<2: continue
        try:
            data.append((float(p[0]),float(p[1])))
        except:
            pass
    return pd.DataFrame(data,columns=["Potential","Current"])

def get_events(y):
    peaks,_=find_peaks(y,distance=PEAK_MIN_DISTANCE,prominence=PEAK_PROMINENCE)
    valleys,_=find_peaks(-np.asarray(y),distance=PEAK_MIN_DISTANCE,prominence=PEAK_PROMINENCE)
    events=[]
    for v in valleys: events.append((v,"valley"))
    for p in peaks: events.append((p,"peak"))
    events.sort(key=lambda x:x[0])
    return peaks,valleys,events

def process_gcd():
    files=[f for f in Path(".").glob("*.txt") if detect_type(read_text(f))=="GCD"]
    if not files:
        print("未发现GCD文件")
        return
    files.sort(key=simplify_name)

    result={}
    durations=[]
    peak_marks={}

    for f in files:
        df=extract_gcd(f)
        y=df["Voltage"].to_numpy()
        t=df["Time"].to_numpy()
        peaks,valleys,events=get_events(y)

        need=max(GCD_END_CYCLE-1,2)
        if len(events)<need:
            print(f"{f.name} 圈数不足")
            continue

        start_idx=events[GCD_START_CYCLE-2][0]
        end_idx=events[GCD_END_CYCLE-1][0]
        seg=y[start_idx:end_idx+1]

        # 计算第3圈时间差（峰到谷）
        if len(peaks)>0 and len(valleys)>1:
            p1=peaks[0]
            v2=valleys[1]
            durations.append([simplify_name(f), t[v2]-t[p1]])

        result[str(simplify_name(f))]=seg
        peak_marks[str(simplify_name(f))]=int(np.argmax(seg))+2

    if not result:
        print("GCD文件处理失败，无数据可写入。")
        return

    max_len=max(len(v) for v in result.values())
    out=pd.DataFrame({"Time":np.arange(max_len)*GCD_TIME_STEP})

    for k,v in result.items():
        arr=np.full(max_len,np.nan)
        arr[:len(v)]=v
        out[k]=arr

    with pd.ExcelWriter(OUTPUT_GCD,engine="openpyxl") as writer:
        out.to_excel(writer,sheet_name="GCD_Data",index=False)
        pd.DataFrame(durations,columns=["Sample","Peak_to_Valley_Time"]).to_excel(
            writer,sheet_name="Duration",index=False
        )

    wb=load_workbook(OUTPUT_GCD)
    ws=wb["GCD_Data"]
    fill=PatternFill("solid",fgColor="FFFF00")

    for c in range(2,ws.max_column+1):
        name=str(ws.cell(1,c).value)
        if name in peak_marks:
            ws.cell(peak_marks[name],c).fill=fill

    wb.save(OUTPUT_GCD)
    print("GCD完成")

def cv_turn_points(potential):
    vmax=np.max(potential)
    vmin=np.min(potential)
    highs=np.where(np.abs(potential-vmax)<=CV_TOLERANCE)[0]
    lows=np.where(np.abs(potential-vmin)<=CV_TOLERANCE)[0]
    pts=[]
    for x in highs: pts.append((x,"H"))
    for x in lows: pts.append((x,"L"))
    pts.sort(key=lambda x:x[0])
    filtered=[]
    for p in pts:
        if not filtered:
            filtered.append(p)
        elif filtered[-1][1]!=p[1]:
            filtered.append(p)
    return filtered

def process_cv():
    files=[f for f in Path(".").glob("*.txt") if detect_type(read_text(f))=="CV"]
    if not files:
        print("未发现CV文件")
        return

    # 用户输入板重
    try:
        mass_mg=float(input("请输入板重(mg): "))
        mass_g=mass_mg/1000.0
        if mass_g<=0:
            raise ValueError("板重必须大于0")
    except Exception as e:
        print("输入错误:",e)
        return

    files.sort(key=simplify_name)

    info=[]
    min_cycles=999999

    for f in files:
        df=extract_cv(f)
        tp=cv_turn_points(df["Potential"].to_numpy())
        cycles=len(tp)-1
        min_cycles=min(min_cycles,cycles)
        info.append((f,df,tp))

    if min_cycles < 4:
        use_start=2
        use_end=3
    else:
        use_start=3
        use_end=4

    out=None
    base_potential=None

    for f,df,tp in info:
        if len(tp) < use_end+1:
            raise ValueError(f"{f.name} 圈数不足")

        start_idx=tp[use_start-1][0]
        end_idx=tp[use_end][0]

        seg=df.iloc[start_idx:end_idx+1].reset_index(drop=True)
        seg["Current"] = seg["Current"]/mass_g  # 归一化

        potential=seg["Potential"].to_numpy()

        if out is None:
            out=pd.DataFrame()
            out["Potential"]=potential
            base_potential=potential
        else:
            if len(potential)!=len(base_potential):
                raise ValueError(f"{f.name} 电压长度不一致")
            if not np.allclose(potential,base_potential,atol=CV_TOLERANCE):
                raise ValueError(f"{f.name} 电压序列不一致")

        out[str(simplify_name(f))]=seg["Current"]

    out.to_excel(OUTPUT_CV,index=False)
    print(f"CV完成，电流已除以板重({mass_mg} mg)归一化, 自动选取第{use_start}+{use_end}圈")

def main():
    try:
        process_gcd()
        process_cv()
    except Exception as e:
        print("错误:",e)

if __name__=="__main__":
    main()
