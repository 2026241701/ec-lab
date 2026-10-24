"""Lightweight file-kind classification for the processing workspace.

Stdlib-only so it can be imported on the GUI thread without pulling in
numpy / pandas / openpyxl. The CDI (conductivity) check lazily imports
openpyxl only when it actually needs to open an .xlsx workbook.

Kinds: CV / GCD / EIS / CDI. ``None`` means "not recognised here".
"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

KINDS = ("CV", "GCD", "EIS", "CDI")
KIND_PRIORITY = {"CV": 0, "GCD": 1, "EIS": 2, "CDI": 3}
KIND_LABELS = {
    "CV": "CV 循环伏安",
    "GCD": "GCD 恒流充放电",
    "EIS": "EIS 交流阻抗",
    "CDI": "CDI 电导率",
    "other": "其他 / 未识别",
}


def detect_kind_text(text):
    """Mirror of eclab.detect_kind for the first chunk of a text file."""
    head = text[:400]
    if "Cyclic Voltammetry" in head:
        return "CV"
    if "Chronopotentiometry" in head:
        return "GCD"
    if "A.C. Impedance" in head or "Freq/Hz" in text[:2000]:
        return "EIS"
    return None


def _compact(value):
    return re.sub(r"[\s_()（）\[\]单位]+", "", str(value or "").strip().lower())


def _header_kind(value):
    compact = _compact(value)
    if any(term in compact for term in ("测量时间", "测试时间", "采集时间", "elapsed", "timestamp")) \
            or compact in {"time", "times", "timemin", "timeh", "datetime", "date/time"} \
            or "时间" in compact:
        return "time"
    if any(term in compact for term in ("检测结果", "电导率", "conductivity")) or compact == "cond":
        return "conductivity"
    return None


def _row_kinds(cells):
    kinds = [_header_kind(v) for v in cells]
    return "time" in kinds and "conductivity" in kinds


def _looks_like_conductivity_text(path):
    """Header / first-lines sniff for CSV, TSV or plain-text conductivity."""
    try:
        sample = path.read_bytes()[:8000]
    except OSError:
        return False
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            content = sample.decode(encoding)
            if encoding == "utf-16" and not sample.startswith((b"\xff\xfe", b"\xfe\xff")):
                continue
            break
        except UnicodeError:
            continue
    else:
        content = sample.decode("utf-8", errors="replace")

    delimiter = "\t" if "\t" in content else "," if "," in content else ";"
    if delimiter:
        reader = csv.reader(io.StringIO(content), delimiter=delimiter)
        for row in reader:
            if _row_kinds(row):
                return True
            if reader.line_num > 30:
                break
    else:
        for line in content.splitlines()[:30]:
            if _row_kinds(re.split(r"\s+", line.strip())):
                return True
    return False


def _looks_like_conductivity_xlsx(path):
    try:
        from openpyxl import load_workbook
    except Exception:
        return False
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return False
    try:
        sheet = book.worksheets[0]
        for _i, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if _i > 30:
                break
            if _row_kinds(row):
                return True
        return False
    finally:
        book.close()


def classify(path):
    """Classify a single file into CV / GCD / EIS / CDI, or return None."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".txt":
        try:
            raw = path.read_bytes()[:4000]
        except OSError:
            return None
        text = None
        for encoding in ("utf-8-sig", "utf-8", "gbk"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = raw.decode("utf-8", errors="replace")
        kind = detect_kind_text(text)
        if kind:
            return kind
        if _looks_like_conductivity_text(path):
            return "CDI"
        return None
    if suffix in (".xlsx", ".csv", ".tsv"):
        if suffix == ".xlsx":
            return "CDI" if _looks_like_conductivity_xlsx(path) else None
        return "CDI" if _looks_like_conductivity_text(path) else None
    return None


def classify_paths(paths):
    """Classify an iterable of paths, returning a list of (path, kind) pairs."""
    return [(str(Path(p)), classify(p)) for p in paths]
